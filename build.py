#!/usr/bin/env python3
"""
Build the Venomous Snakes Explained website from the workbook.

    python build.py                      # build into dist/ (fetches missing photos)
    python build.py --no-photos          # build without contacting iNaturalist
    python build.py --refresh-photos     # re-check every species on iNaturalist
    python build.py --base-url https://your-domain.example   # also write sitemap.xml

Input : data/snake_venom_danger.xlsx  (save it in Excel after editing, so formula results are stored)
Output: dist/  - upload this folder to Cloudflare Pages.
"""
import argparse, csv, html, json, math, re, shutil, sys, time, urllib.parse, urllib.request, warnings
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is missing. Install it with:  pip install openpyxl")
warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parent
SRC, DIST, PHOTOS = ROOT / "src", ROOT / "dist", ROOT / "photos"
WORKBOOK = ROOT / "data" / "snake_venom_danger.xlsx"
SITE = "Venomous Snakes Explained"
ALLOWED_LICENCES = {"cc0", "cc-by", "cc-by-sa"}      # no NC/ND: the site may link to a paid book
LICENCE_LABEL = {"cc0": "CC0", "cc-by": "CC BY", "cc-by-sa": "CC BY-SA"}
INAT = "https://api.inaturalist.org/v1"
UA = "VenomousSnakesExplained-site-build/1.0 (static site; one request per second)"

E = lambda s: html.escape(str(s)) if s is not None else ""
isnum = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
def slug(sci): return "sp-" + re.sub(r"[^a-z0-9]+", "-", sci.lower()).strip("-") + ".html"
def stem(sci): return slug(sci)[3:-5]

# ----------------------------------------------------------------------------- workbook
def load(path):
    w = openpyxl.load_workbook(path, data_only=True)
    need = ["Species", "Scores", "Combined", "LD50 Sources", "Yield Sources", "Clinical Sources", "Settings"]
    missing = [n for n in need if n not in w.sheetnames]
    if missing:
        sys.exit(f"Workbook is missing sheets: {missing}")
    rows = lambda sh: [r for r in w[sh].iter_rows(min_row=2, values_only=True) if r[0]]
    sp = {r[0]: r for r in rows("Species")}
    sc = {r[0]: r for r in rows("Scores")}
    cb = {r[0]: r for r in rows("Combined")}
    if all(sc[n][2] is None for n in sc):
        sys.exit("The workbook has no stored formula results. Open it in Excel, save it, and run the build again.")
    group = lambda sh: {k: [r for r in rows(sh) if r[0] == k] for k in sp}
    st = w["Settings"]
    return dict(sp=sp, sc=sc, cb=cb, ld=group("LD50 Sources"), y=group("Yield Sources"), c=group("Clinical Sources"),
                bite_scale={st.cell(r, 2).value: st.cell(r, 1).value for r in range(50, 55)},
                anchors=dict(ldMost=round(st["B22"].value, 4), ldLeast=round(st["B23"].value, 3), yMin=st["B24"].value, yMax=st["B25"].value))

# ----------------------------------------------------------------------------- public wording
def public(t):
    """Turn workbook shorthand into wording fit for the public site."""
    if not t: return t
    t = str(t)
    t = t.replace("Not recorded in v3.2; needs a source", "Source not yet found")
    t = re.sub(r"\bTier ([A-D])\b", r"grade \1", t)
    t = re.sub(r"\s*\[Added [^\]]*\]", "", t)
    t = re.sub(r"\s*\([^()]*20\d\d-\d\d-\d\d[^()]*\)", "", t)
    t = t.replace("this row is", "this page covers")
    t = re.sub(r"\s*\*$", " (source not yet found)", t.strip())
    return t.strip()

def per_year(t):
    return re.sub(r"^([<~]?[\d.,\-]+)", r"\1 a year", t) if t else t

def public_flags(f):
    if not f: return ""
    out = []
    for p in str(f).split("; "):
        m = re.match(r"(LD50|yield) sources differ ([\d.]+)x", p)
        if m: out.append(f"{m.group(1).capitalize() if m.group(1) == 'yield' else 'LD50'} sources differ {m.group(2)}×"); continue
        m = re.match(r"LD50 route (\w+)", p)
        if m: out.append("LD50 route not recorded" if m.group(1) == "unknown" else f"LD50 not tested under the skin (route: {m.group(1)})"); continue
        m = re.match(r"(LD50|yield) Tier D", p)
        if m: out.append(("LD50" if m.group(1) == "LD50" else "Venom yield") + " has no traceable source"); continue
        out.append(p)
    return "; ".join(out)

KEYS = ["Clinical", "Antivenom", "Taxonomy", "Venom", "Onset", "Bite propensity", "COVERAGE GAP", "LD50", "Encounter", "Sources"]
def parse_notes(t):
    out = {}
    if not t: return out
    t = re.sub(r"\s*\[Added [^\]]*\]", "", t)
    t = re.sub(r"\s*\([^()]*20\d\d-\d\d-\d\d[^()]*\)", "", t)
    t = re.sub(r"\s*Rated [^.]*20\d\d-\d\d-\d\d\.", "", t)
    t = t.replace("Antivenom: none - ", "Antivenom: None. ")
    pat = re.compile(r"(?:(?<=^)|(?<=\.\s)|(?<=\)\s))(" + "|".join(re.escape(k) for k in KEYS) + r")\b[^:.]{0,60}:\s*")
    ms = list(pat.finditer(t))
    for i, m in enumerate(ms):
        end = ms[i + 1].start() if i + 1 < len(ms) else len(t)
        seg = t[m.end():end].strip()
        # drop sentences that only make sense inside the workbook
        seg = " ".join(s for s in re.split(r"(?<=\.)\s+", seg) if "v3.2" not in s and " row" not in s)
        if seg: out.setdefault(m.group(1), []).append(public(seg))
    return out

# ----------------------------------------------------------------------------- photos
def http_json(url):
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "application/json"})
    with urllib.request.urlopen(req, timeout=25) as r:
        return json.load(r)

def pick_photo(taxon):
    """First curated taxon photo with an allowed licence, default photo first."""
    cands = []
    if taxon.get("default_photo"): cands.append(taxon["default_photo"])
    cands += [tp.get("photo", {}) for tp in taxon.get("taxon_photos", []) or []]
    for p in cands:
        if (p.get("license_code") or "").lower() in ALLOWED_LICENCES and (p.get("medium_url") or p.get("url")):
            return p
    return None

def find_photo(sci, getter=http_json, pause=1.1):
    q = urllib.parse.quote(sci)
    res = getter(f"{INAT}/taxa?q={q}&rank=species&is_active=true&per_page=10").get("results", [])
    time.sleep(pause)
    match = next((t for t in res if (t.get("name") or "").lower() == sci.lower()), None) \
        or next((t for t in res if (t.get("matched_term") or "").lower() == sci.lower()), None)
    if not match:
        return {"status": "no iNaturalist match"}
    full = getter(f"{INAT}/taxa/{match['id']}").get("results", [match])[0]
    time.sleep(pause)
    p = pick_photo(full)
    if not p:
        return {"status": "no photo with an allowed licence", "inat_name": full.get("name"), "taxon_id": full.get("id")}
    url = p.get("medium_url") or p.get("url")
    url = re.sub(r"/(square|small|medium|thumb)\.", "/large.", url)
    return {"status": "ok", "inat_name": full.get("name"), "taxon_id": full.get("id"), "photo_id": p.get("id"),
            "url": url, "licence": p.get("license_code").lower(), "attribution": p.get("attribution"),
            "page": f"https://www.inaturalist.org/photos/{p.get('id')}"}

def load_overrides():
    """photos/overrides.csv: scientific_name,url,attribution,licence,page  (licence: cc0 / cc-by / cc-by-sa, or 'none' to show no photo)"""
    f = PHOTOS / "overrides.csv"
    if not f.exists(): return {}
    with f.open(newline="", encoding="utf-8") as fh:
        return {r["scientific_name"].strip(): r for r in csv.DictReader(fh) if r.get("scientific_name")}

def photos(species, fetch, refresh):
    cache_f = PHOTOS / "photos.json"
    cache = json.loads(cache_f.read_text()) if cache_f.exists() else {}
    overrides = load_overrides()
    img_dir = PHOTOS / "images"; img_dir.mkdir(parents=True, exist_ok=True)
    online = fetch
    for sci in species:
        if sci in overrides:
            o = overrides[sci]
            if (o.get("licence") or "").lower() == "none":
                cache[sci] = {"status": "no photo (override)"}; continue
            cache[sci] = {"status": "ok", "url": o["url"], "licence": o["licence"].lower(), "attribution": o["attribution"], "page": o.get("page", ""), "override": True}
        elif online and (refresh or sci not in cache or cache[sci].get("status") not in ("ok",)):
            try:
                cache[sci] = find_photo(sci)
                print(f"  photo  {sci}: {cache[sci]['status']}")
            except Exception as e:
                print(f"  photo  {sci}: could not reach iNaturalist ({e}); continuing without photos")
                online = False
        c = cache.get(sci, {})
        img = img_dir / (stem(sci) + ".jpg")
        if c.get("status") == "ok" and online and (refresh or not img.exists()):
            try:
                req = urllib.request.Request(c["url"], headers={"User-Agent": UA})
                with urllib.request.urlopen(req, timeout=40) as r: img.write_bytes(r.read())
                time.sleep(0.5)
            except Exception as e:
                print(f"  photo  {sci}: download failed ({e})")
        c["file"] = img.name if img.exists() and c.get("status") == "ok" else None
        cache[sci] = c
    cache_f.write_text(json.dumps(cache, indent=1, ensure_ascii=False))
    return cache

# ----------------------------------------------------------------------------- species pages
FAMILY = {"Elapidae": "Elapid", "Viperidae": "Viper", "Colubridae": "Rear-fanged colubrid"}
GRADE = {
    "A": "Built on named, peer-reviewed studies with the method stated.",
    "B": "Built on expert databases or WHO documents, with the route stated.",
    "C": "Built on secondary compilations rather than the original lab studies. The ranking is unlikely to change much, but the exact figures could.",
    "D": "At least one input has no traceable source yet. Treat the Venom Hazard as provisional.",
}
EXTRA_FACTS = {
    "Daboia russelii": ("Bites", '43% of bites in Indian studies that identified the snake <span class="tier A" title="Grade A">A</span>',
                        ("Share of identified snakebites, India", "43%", "Suraweera et al. (2020), <i>eLife</i> 9:e54076, review of 87,590 bites. The same study estimates about 58,000 snakebite deaths a year in India from all species. No study splits those deaths by species.", "A")),
}

def fmt_num(v):
    if not isnum(v): return E(v)
    return f"{v:,.4f}".rstrip("0").rstrip(".") if v < 1 else f"{v:,.1f}".rstrip("0").rstrip(".")

def dim(label, v, text):
    val = f'<span class="num">{v:.1f}</span>' if isnum(v) else '<span class="num" style="color:var(--muted)">–</span>'
    w = max(0, min(100, v)) if isnum(v) else 0
    return f'<div class="dim"><div class="row"><span>{label}</span>{val}</div><div class="track"><div class="fill" style="width:{w:.1f}%"></div></div><p>{text}</p></div>'

def card(cls, title, q, v, rank, n, body, missing=None):
    big = f'<span class="big">{v:.1f}</span>' if isnum(v) else '<span class="big" style="font-size:1.3rem;color:var(--muted)">Not scored</span>'
    rk = f'<div class="rank">Rank<br><b>{rank}</b> of {n}</div>' if isnum(v) and rank else ""
    miss = f'<p class="caption">{E(missing)}</p>' if missing else ""
    return f'<article class="score {cls}"><div class="head"><div style="display:grid;gap:2px"><span class="eyebrow">{title}</span><span style="font-size:0.88rem;color:var(--ink-2)">{q}</span>{big}</div>{rk}</div>{miss}{body}</article>'

HEAD = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<title>{title}</title>
<meta name="description" content="{desc}">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=IBM+Plex+Mono:wght@400;500&family=Public+Sans:wght@400;600;700&family=Zilla+Slab:ital,wght@0,500;0,600;0,700;1,500&display=swap">
<link rel="stylesheet" href="site.css">
</head>
<body data-page="{page}">
<script src="common.js"></script>
"""

def species_pages(D, photo_info):
    sp, sc, cb = D["sp"], D["sc"], D["cb"]
    nVH = sum(isnum(s[2]) for s in sc.values()); nHR = sum(isnum(s[4]) for s in sc.values()); nC = sum(isnum(c[11]) for c in cb.values())
    pts = [(n, sc[n][2], sc[n][4]) for n in sp if isnum(sc[n][2]) and isnum(sc[n][4])]
    for n, r in sp.items():
        s, c = sc[n], cb[n]
        common, other, fam, region, effects = r[1], r[2], r[3], r[4], r[5]
        ld, route, ymax = r[6], r[7], r[11]
        enc_raw, bite_raw, onset, access, deaths, cfr = r[14], r[15], r[16], r[17], r[18], r[19]
        reason = r[30] if len(r) > 30 else None
        notes = parse_notes(r[20])
        vh, hr, comb = s[2], s[4], c[11]
        facts = [("Venom", (E(effects) + ' <a href="venom-types.html" style="font-size:0.85rem">what this means</a>') if effects else "Not recorded yet"),
                 ("Onset", (E(onset) + ", if untreated") if onset else "Not recorded yet"),
                 ("Antivenom", E(access) if access else "Not recorded yet"),
                 ("Deaths", E(per_year(public(deaths))) if deaths else "No sourced figure yet")]
        if cfr: facts.append(("Case fatality", E(public(cfr))))
        if n in EXTRA_FACTS: facts.append(EXTRA_FACTS[n][:2])
        parts = []
        if isnum(vh): parts.append(f"{s[3]} of {nVH} for Venom Hazard")
        if isnum(hr): parts.append(f"{s[5]} of {nHR} for Human Risk")
        summary = f"{E(common)} " + ("ranks " + " and ".join(parts) + "." if parts else "has no scores yet because key data is missing.")

        pot_txt = f"LD50 of {fmt_num(ld)} mg/kg ({E(route)} injection)." if isnum(ld) else "No LD50 figure yet."
        qty_txt = f"Up to {fmt_num(ymax)} mg of venom in a single milking." if isnum(ymax) else "No venom yield figure yet, so no Venom Hazard."
        vh_card = card("", "Venom Hazard", "How dangerous is the venom itself?", vh, s[3], nVH,
                       dim("Potency", s[10], pot_txt) + dim("Quantity", s[11], qty_txt), None if isnum(vh) else f"Missing: {vh}.")
        enc_txt = (f"Rated {enc_raw} out of 10. " if isnum(enc_raw) else "Not rated yet. ") + E(reason or "")
        bite_txt = f"Level {bite_raw} of 5: {E(D['bite_scale'].get(bite_raw, ''))}." if isnum(bite_raw) else "Not scored: temperament not documented."
        ons_txt = E(onset) + ". Slow onset scores higher because people often delay treatment." if onset else "Not recorded yet."
        acc_txt = E(access) + "." if access else "Not recorded yet."
        hr_card = card("", "Human Risk", "How likely is a bite to happen and go badly?", hr, s[5], nHR,
                       dim("Encounter", s[12], enc_txt) + dim("Bite propensity", s[13], bite_txt) + dim("Onset", s[14], ons_txt) + dim("Antivenom access", s[15], acc_txt),
                       None if isnum(hr) else "Fewer than three of the four factors are known.")
        comb_card = card("comb", "Combined", "Both questions averaged", comb, c[12], nC,
                         '<p style="font-size:0.88rem;color:var(--ink-2)">An equal average of Venom Hazard and Human Risk. Always read it with the two scores beside it.</p>',
                         None if isnum(comb) else f"Missing: {comb}.")

        bite_does = " ".join(notes.get("Clinical", [])) or "No clinical summary recorded yet."
        extra = "".join(f"<p><b>{label}.</b> {E(t)}</p>" for k, label in (("Venom", "Venom"), ("Antivenom", "Antivenom source"), ("Taxonomy", "Name and taxonomy")) for t in notes.get(k, []))

        rows = []
        for x in D["ld"][n]:
            if x[8] == "Used": rows.append(("LD50 (mice)", f"{fmt_num(x[2])} mg/kg · {E(x[3])}", E(public(x[6])) + ("; route assumed" if x[4] == "assumed" else ""), x[7]))
        for x in D["y"][n]:
            if x[10] == "Used":
                rng = f"{fmt_num(x[2])}–{fmt_num(x[3])} mg" if isnum(x[2]) and x[2] != x[3] else f"{fmt_num(x[3])} mg"
                rows.append(("Venom yield", rng, E(public(x[8])) + (f"; {E(public(x[4]))}" if x[4] else ""), x[9]))
        for x in D["c"][n]:
            if x[10] == "Used": rows.append((E(x[2]), E(per_year(public(x[12] or "")) if x[2] == "Annual deaths" else public(x[12] or "")), E(public(x[8])), x[9]))
        if not deaths: rows.append(("Deaths a year", "–", "No sourced figure yet.", None))
        if n in EXTRA_FACTS: rows.append(EXTRA_FACTS[n][2])
        tierb = lambda t: f'<span class="tier {t}">{t}</span>' if t else ""
        src_html = "".join(f'<tr><td>{a}</td><td class="r num">{b}</td><td>{cc}</td><td>{tierb(t)}</td></tr>' for a, b, cc, t in rows)
        others = [f"LD50 {fmt_num(x[2])} mg/kg ({E(x[3])}) · {E(x[8]).lower()} · grade {E(x[7])}" for x in D["ld"][n] if x[8] != "Used"] + \
                 [f"Yield {fmt_num(x[2])}–{fmt_num(x[3])} mg · {E(x[10]).lower()} · grade {E(x[9])}" for x in D["y"][n] if x[10] != "Used"]
        others_html = (f'<details><summary style="cursor:pointer;font-weight:600;font-size:0.9rem">Other figures on record ({len(others)})</summary><ul style="margin:8px 0 0;padding-left:1.1em;font-size:0.88rem;display:grid;gap:4px">'
                       + "".join(f"<li>{o}</li>" for o in others) + "</ul></details>") if others else ""

        grade, flags = s[6], public_flags(s[7])
        flag_p = f"<p><b>Flags:</b> {E(flags)}.</p>" if flags else ""
        if grade in GRADE:
            grade_html = f'<section class="grade"><span class="tier {grade}">{grade}</span><div style="display:grid;gap:6px"><h3>How sure are we?</h3><p>{GRADE[grade]}</p>{flag_p}</div></section>'
        else:
            grade_html = f'<section class="grade"><span class="tier D" style="color:var(--muted)">–</span><div style="display:grid;gap:6px"><h3>How sure are we?</h3><p>No Venom Hazard yet, so no data grade.</p>{flag_p}</div></section>'

        sim_html = ""
        if isnum(vh) and isnum(hr):
            near = sorted((math.hypot(a - vh, b - hr), m) for m, a, b in pts if m != n)[:5]
            sim = "".join(f'<a class="chip" href="{slug(m)}">{E(sp[m][1])} <span class="num" style="color:var(--muted)">{sc[m][2]:.0f} / {sc[m][4]:.0f}</span></a>' for _, m in near)
            sim_html = f'<section><h2>Similar species</h2><p class="caption">Closest in both Venom Hazard and Human Risk.</p><div>{sim}</div></section>'

        ph = photo_info.get(n, {})
        if ph.get("file"):
            photo_html = (f'<figure class="photo-fig"><img class="photo-img" src="photos/{ph["file"]}" alt="{E(common)} ({E(n)})" loading="eager">'
                          f'<figcaption class="caption">Photo: {E(ph.get("attribution", ""))}, via <a href="{E(ph.get("page", ""))}" rel="noopener">iNaturalist</a></figcaption></figure>')
        else:
            photo_html = '<figure class="photo-fig"><div class="photo photo-none" role="img" aria-label="No photo yet">Photo coming soon</div></figure>'
        aka = f'<span class="aka">Also called {E(other).replace("; ", ", ")}</span>' if other else ""

        page = HEAD.format(title=f"{E(common)} | {SITE}", desc=f"{E(common)} ({E(n)}): how dangerous its venom is, how likely a bite is to go badly, and where every number comes from.", page="species.html") + f"""<main class="species-page">
  <nav class="crumbs" aria-label="Breadcrumb"><a href="species.html">All species</a> / {E(common)}</nav>
  <section class="top">
    <div class="ident">
      <span class="eyebrow">{E(FAMILY.get(fam, fam))} · {E(region)}</span>
      <h1>{E(common)}</h1>
      <span class="sci">{E(n)}</span>
      {aka}
      <p class="summary">{summary}</p>
      <dl class="facts">{"".join(f"<dt>{k}</dt><dd>{v}</dd>" for k, v in facts)}</dl>
    </div>
    {photo_html}
  </section>
  <section>
    <h2>How it scores</h2>
    <div class="scores">{vh_card}{hr_card}{comb_card}</div>
    <p class="caption">Scores are relative to the species on this site. Bars run from 0 to 100. <a href="scores.html">How these scores are built</a>.</p>
  </section>
  <section class="with-notes">
    <div class="col">
      <h2>What a bite does</h2>
      <p>{E(bite_does)}</p>
      {extra}
    </div>
    <aside class="note"><b>If you are bitten</b><p>Get to a hospital that holds antivenom. Keep still, remove rings and watches, and do not cut, suck or tie off the bite. Use the red “Bitten?” button at the top of any page.</p></aside>
  </section>
  <section>
    <h2>The numbers and where they come from</h2>
    <div class="table-wrap"><table><thead><tr><th>Measure</th><th class="r">Value</th><th>Source</th><th>Grade</th></tr></thead><tbody>{src_html}</tbody></table></div>
    {others_html}
  </section>
  {grade_html}
  {sim_html}
</main>
</body>
</html>
"""
        (DIST / slug(n)).write_text(page, encoding="utf-8")

# ----------------------------------------------------------------------------- data.js and static pages
def data_js(D):
    out = []
    num = lambda v: round(v, 1) if isnum(v) else v
    for n, r in D["sp"].items():
        s, c = D["sc"][n], D["cb"][n]
        out.append(dict(sci=n, common=r[1], other=r[2], family=r[3], region=r[4], effects=r[5], ld50=r[6], route=r[7], ldTier=r[8],
                        yieldMax=r[11], yTier=r[12], onset=r[16], access=r[17], deaths=public(r[18]), cfr=public(r[19]),
                        vh=num(s[2]), hr=num(s[4]), comb=num(c[11]), pot=num(s[10]), qty=num(s[11]), enc=num(s[12]), bite=num(s[13]),
                        ons=num(s[14]), acc=num(s[15]), dims=s[16], grade=s[6], flags=public_flags(s[7]),
                        vhRank=s[3], hrRank=s[5], combRank=c[12]))
    (DIST / "data.js").write_text("window.SPECIES=" + json.dumps(out, ensure_ascii=False, separators=(",", ":")) +
                                  ";\nwindow.ANCHORS=" + json.dumps(D["anchors"]) + ";\n", encoding="utf-8")

TITLES = {"scores.html": ("Scores explained", "How Venom Hazard and Human Risk are built: LD50, venom yield, the four risk factors, and how far to trust each number."),
          "venom-types.html": ("Venom types and combinations", "What neurotoxic, haemotoxic, cytotoxic and myotoxic venoms do to the body, and which snakes combine them."),
          "species.html": ("All species", "All 94 venomous snakes on this site, with both scores. Search, filter and sort.")}

def static_pages():
    for f in SRC.iterdir():
        if f.suffix == ".html":
            t = f.read_text(encoding="utf-8")
            t = re.sub(r'\s*<div class="editor">.*?</div>', "", t, flags=re.S)   # review notes never ship
            if f.name in TITLES:
                title, desc = TITLES[f.name]
                t = re.sub(r"<title>.*?</title>", f"<title>{title} | {SITE}</title>", t, count=1)
                if 'name="description"' not in t:
                    t = t.replace("</title>", f'</title>\n<meta name="description" content="{E(desc)}">', 1)
            (DIST / f.name).write_text(t, encoding="utf-8")
        elif f.suffix in (".css", ".js"):
            shutil.copy(f, DIST / f.name)

def extras(D, base_url):
    (DIST / "404.html").write_text(HEAD.format(title=f"Page not found | {SITE}", desc="Page not found.", page="") +
        '<main><section class="col"><span class="eyebrow">404</span><h1>Page not found</h1><p class="lede">That page doesn’t exist. Try the <a href="species.html">species list</a> or the <a href="index.html">home page</a>.</p></section></main>\n</body>\n</html>\n', encoding="utf-8")
    (DIST / "_headers").write_text("/*\n  X-Content-Type-Options: nosniff\n  Referrer-Policy: strict-origin-when-cross-origin\n  X-Frame-Options: SAMEORIGIN\n/photos/*\n  Cache-Control: public, max-age=2592000\n", encoding="utf-8")
    lines = ["User-agent: *", "Allow: /"]
    if base_url:
        base = base_url.rstrip("/")
        urls = ["", "scores", "venom-types", "species"] + [slug(n)[:-5] for n in D["sp"]]
        (DIST / "sitemap.xml").write_text('<?xml version="1.0" encoding="UTF-8"?>\n<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n' +
            "".join(f"  <url><loc>{base}/{u}</loc></url>\n" for u in urls) + "</urlset>\n", encoding="utf-8")
        lines.append(f"Sitemap: {base}/sitemap.xml")
    (DIST / "robots.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")

# ----------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--no-photos", action="store_true", help="do not contact iNaturalist (uses photos already downloaded)")
    ap.add_argument("--refresh-photos", action="store_true", help="re-check every species on iNaturalist")
    ap.add_argument("--base-url", default="", help="site address, e.g. https://example.com, to write sitemap.xml")
    ap.add_argument("--workbook", default=str(WORKBOOK))
    a = ap.parse_args()

    print(f"Reading {a.workbook}")
    D = load(a.workbook)
    if DIST.exists(): shutil.rmtree(DIST)
    DIST.mkdir()
    print(f"Photos ({'offline' if a.no_photos else 'checking iNaturalist'})")
    info = photos(list(D["sp"]), fetch=not a.no_photos, refresh=a.refresh_photos)
    (DIST / "photos").mkdir()
    for n, c in info.items():
        if c.get("file"): shutil.copy(PHOTOS / "images" / c["file"], DIST / "photos" / c["file"])
    data_js(D)
    static_pages()
    species_pages(D, info)
    extras(D, a.base_url)
    have = sum(1 for n in D["sp"] if info.get(n, {}).get("file"))
    missing = [n for n in D["sp"] if not info.get(n, {}).get("file")]
    print(f"\nBuilt {len(list(DIST.glob('*.html')))} pages into {DIST}")
    print(f"Photos: {have} of {len(D['sp'])} species")
    if missing:
        (ROOT / "photos" / "missing.txt").write_text("\n".join(f"{n}\t{info.get(n, {}).get('status', 'not checked')}" for n in missing) + "\n")
        print(f"  {len(missing)} without a photo - see photos/missing.txt; add them to photos/overrides.csv if you find one")

if __name__ == "__main__":
    main()
